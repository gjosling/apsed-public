#!/usr/bin/env python3
"""
Parse agency-level total headcount from Table 2 of all APSED Excel releases.

Each release contains "Table 2: All employees: agency by employment category"
with current and prior year headcounts per agency, broken down into Ongoing,
Non-ongoing, and Total. This script extracts only the current-snapshot columns.

Portfolio department rows record only the core department's direct employees.
Sub-agencies within the same portfolio are listed as separate rows with their
own headcounts. Summing all rows (portfolio departments and sub-agencies alike)
gives the complete APS headcount.

Outputs data/apsed_agency_headcount.csv:
  snapshot_date, agency, agency_canonical, is_sub_agency, parent_agency,
  headcount_ongoing, headcount_non_ongoing, headcount_total

Usage:
    uv run python pipeline/04_parse_agency_headcount.py
    uv run python pipeline/04_parse_agency_headcount.py 2024-06-30
"""
import csv
import datetime
import sys
import traceback
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from utils import XLSX_DIR, FOOTER_PREFIXES, _AGENCY_CANONICAL, snapshot_date_from_path

OUTPUT = Path("data/apsed_agency_headcount.csv")

OUTPUT_COLS = [
    "snapshot_date",
    "agency",
    "agency_canonical",
    "is_sub_agency",
    "parent_agency",
    "headcount_ongoing",
    "headcount_non_ongoing",
    "headcount_total",
]

_SCHEMA = pa.schema([
    ("snapshot_date",         pa.date32()),
    ("agency",                pa.string()),
    ("agency_canonical",      pa.string()),
    ("is_sub_agency",         pa.bool_()),
    ("parent_agency",         pa.string()),
    ("headcount_ongoing",     pa.int32()),
    ("headcount_non_ongoing", pa.int32()),
    ("headcount_total",       pa.int32()),
])

# ── Column detection ──────────────────────────────────────────────────────────

def find_current_year_cols(ws, snapshot_year: int) -> tuple[int, int, int]:
    """Return (ongoing_col, non_ongoing_col, total_col) for the current snapshot year.

    Table 2 has a side-by-side two-year layout. Row 3 (0-based) carries year
    values — either plain integers or datetime objects — repeated once per
    column group (Ongoing, Non-ongoing, Total). The three column indices
    matching snapshot_year are the current-year ongoing, non-ongoing, and total
    columns in that order.
    """
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        if i > 5:
            break
        matches = []
        for j, v in enumerate(row):
            if v is None or j == 0:
                continue
            yr = v.year if hasattr(v, "year") else None
            if yr is None:
                try:
                    yr = int(float(str(v)))
                except (ValueError, TypeError):
                    pass
            if yr == snapshot_year:
                matches.append(j)
        if len(matches) == 3:
            return matches[0], matches[1], matches[2]

    raise ValueError(f"Could not locate current-year columns for {snapshot_year}")

# ── Value parsing ─────────────────────────────────────────────────────────────

def parse_int(v, label: str, context: str) -> int | None:
    if v is None:
        print(f"  [warn] empty {label} cell for {context!r}")
        return None
    try:
        return int(float(str(v).replace(",", "")))
    except (ValueError, TypeError):
        print(f"  [warn] non-numeric {label} value {v!r} for {context!r}")
        return None

# ── Sheet parsing ─────────────────────────────────────────────────────────────

def parse_sheet(ws, snapshot_date: str, snapshot_year: int) -> list[dict]:
    ongoing_col, non_ongoing_col, total_col = find_current_year_cols(ws, snapshot_year)

    records: list[dict] = []
    current_parent: str | None = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= 3:
            continue

        raw = row[0]
        if raw is None:
            continue
        name = str(raw).strip()
        if not name:
            continue

        lower = name.lower()
        if lower == "total" or any(lower.startswith(p) for p in FOOTER_PREFIXES):
            continue
        if name.startswith("*"):  # MoG footnotes in column A
            continue

        is_sub = name.startswith("- ")
        agency_clean = name[2:].strip() if is_sub else name

        if agency_clean.endswith("(NC)"):  # non-current (defunct) agencies
            continue

        if not is_sub:
            current_parent = agency_clean
            parent = None
        else:
            parent = current_parent

        ctx = f"{snapshot_date} / {agency_clean}"
        records.append({
            "snapshot_date":         snapshot_date,
            "agency":                agency_clean,
            "agency_canonical":      _AGENCY_CANONICAL.get(agency_clean, agency_clean),
            "is_sub_agency":         is_sub,
            "parent_agency":         parent,
            "headcount_ongoing":     parse_int(
                                         row[ongoing_col] if ongoing_col < len(row) else None,
                                         "ongoing", ctx),
            "headcount_non_ongoing": parse_int(
                                         row[non_ongoing_col] if non_ongoing_col < len(row) else None,
                                         "non_ongoing", ctx),
            "headcount_total":       parse_int(
                                         row[total_col] if total_col < len(row) else None,
                                         "total", ctx),
        })

    return records

# ── Main ──────────────────────────────────────────────────────────────────────

def main(target_date: str | None = None) -> None:
    files = sorted(XLSX_DIR.glob("apsed_*.xlsx"))
    all_records: list[dict] = []
    processed: list[str] = []

    for path in files:
        date = snapshot_date_from_path(path)
        if not date:
            continue
        if target_date and date != target_date:
            continue

        year = int(date[:4])
        print(f"  [{date}] Opening ...", end=" ", flush=True)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        if "Table 2" not in wb.sheetnames:
            wb.close()
            print("Table 2 not found — skipped")
            continue

        ws = wb["Table 2"]

        try:
            records = parse_sheet(ws, date, year)
        except Exception as exc:
            wb.close()
            print(f"ERROR: {exc}")
            traceback.print_exc()
            continue

        wb.close()

        n_top = sum(1 for r in records if not r["is_sub_agency"])
        n_sub = sum(1 for r in records if r["is_sub_agency"])
        aps_total = sum(
            r["headcount_total"] for r in records if r["headcount_total"] is not None
        )
        print(
            f"{len(records):,} rows "
            f"({n_top} portfolio depts, {n_sub} sub-agencies), "
            f"APS total = {aps_total:,}"
        )
        all_records.extend(records)
        processed.append(date)

    if not all_records:
        print("No rows parsed — nothing written.")
        return

    with open(OUTPUT, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLS)
        writer.writeheader()
        writer.writerows(all_records)

    parquet_path = OUTPUT.with_suffix(".parquet")
    arrays = []
    for col in OUTPUT_COLS:
        field_type = _SCHEMA.field(col).type
        values = [r[col] for r in all_records]
        if field_type == pa.date32():
            values = [datetime.date.fromisoformat(v) for v in values]
        arrays.append(pa.array(values, type=field_type))
    pq.write_table(pa.table(dict(zip(OUTPUT_COLS, arrays)), schema=_SCHEMA), parquet_path)

    print(f"\nWrote {len(all_records):,} rows → {OUTPUT}")
    print(f"                          → {parquet_path}")
    print(f"Dates: {processed}")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    main(target)
