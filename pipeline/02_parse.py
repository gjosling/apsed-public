#!/usr/bin/env python3
"""
Parse the agency × job family table from APSED Excel releases.

Only processes releases that contain a Table 88-equivalent sheet (Jun 2024 onwards).
Outputs long-format CSV: one row per agency × job family × snapshot_date.

Usage:
    uv run python pipeline/02_parse.py                # all eligible files
    uv run python pipeline/02_parse.py 2025-12-31     # single date
"""
import csv
import datetime
import sys
import traceback
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from utils import XLSX_DIR, REGIME_MAP, FOOTER_PREFIXES, _AGENCY_CANONICAL, snapshot_date_from_path

OUTPUT = Path("data/apsed_agency_jf.csv")

# REGIME_MAP is imported from utils.py.

OUTPUT_COLS = [
    "snapshot_date",
    "taxonomy_regime",
    "agency",
    "agency_canonical",
    "is_sub_agency",
    "parent_agency",
    "job_family",
    "job_family_key",
    "headcount",
]

# FOOTER_PREFIXES is imported from utils.py.

# Unclassified-employee column label varies across releases; normalise all to "No data".
#   "No Data"   — Jun 2024, Dec 2024
#   "Zero Data" — Jun 2025
#   "No data"   — Dec 2025, Jun 2026 (Regime 3 canonical)
_FAMILY_NORMALISE: dict[str, str] = {
    "No Data": "No data",
    "Zero Data": "No data",
}

# _AGENCY_CANONICAL is imported from utils.py.

# Maps normalised job family labels to SCREAMING_SNAKE_CASE keys matching the
# gazette dataset's taxonomy. None where no gazette equivalent exists.
_JF_KEY: dict[str, str | None] = {
    "Accounting and Finance":                   "ACCOUNTING_AND_FINANCE",
    "Administration":                           "BUSINESS_AND_ORGANISATIONAL_MANAGEMENT",
    "Business and Organisational Management":   "BUSINESS_AND_ORGANISATIONAL_MANAGEMENT",
    "Communications and Engagement":            "COMMUNICATIONS_AND_ENGAGEMENT",
    "Communications and Marketing":             "COMMUNICATIONS_AND_ENGAGEMENT",
    "Compliance and Regulation":                "COMPLIANCE_AND_REGULATION",
    "Data and Research":                        "DATA_AND_RESEARCH",
    "Engineering and Technical":                "ENGINEERING_AND_TECHNICAL",
    "Health":                                   "HEALTH",
    "Human Resources":                          "HUMAN_RESOURCES",
    "ICT and Digital":                          "ICT_AND_DIGITAL",
    "ICT and Digital Solutions":                "ICT_AND_DIGITAL",
    "Information and Knowledge Management":     "INTELLIGENCE_AND_INFORMATION_MANAGEMENT",
    "Intelligence":                             "INTELLIGENCE_AND_INFORMATION_MANAGEMENT",
    "Intelligence and Information Management":  "INTELLIGENCE_AND_INFORMATION_MANAGEMENT",
    "Legal and Parliamentary":                  "LEGAL_AND_PARLIAMENTARY",
    "Monitoring and Audit":                     None,  # split-absorbs across Accounting and Finance + Compliance in Regime 3
    "No data":                                  None,
    "Policy":                                   "POLICY",
    "Portfolio, Program and Project Management":  "PORTFOLIO_PROGRAM_AND_PROJECT_MANAGEMENT",
    "Science":                                  "SCIENCE",
    "Science and Health":                       None,  # not separable in 2a/2b/2b+
    "Senior Executive":                         None,
    "Service Delivery":                         "SERVICE_DELIVERY",
    "Trades and Labour":                        "TRADES_AND_LABOUR",
}

# Explicit PyArrow schema so snapshot_date is stored as date32, not string.
_SCHEMA = pa.schema([
    ("snapshot_date",    pa.date32()),
    ("taxonomy_regime",  pa.string()),
    ("agency",           pa.string()),
    ("agency_canonical", pa.string()),
    ("is_sub_agency",    pa.bool_()),
    ("parent_agency",    pa.string()),
    ("job_family",       pa.string()),
    ("job_family_key",   pa.string()),
    ("headcount",        pa.int32()),
])

# ── Sheet detection ───────────────────────────────────────────────────────────

def find_agency_jf_sheet(wb: openpyxl.Workbook):
    """Return the worksheet whose title contains "agency by job family", or None."""
    for sname in wb.sheetnames:
        ws = wb[sname]
        for rnum, row in enumerate(ws.iter_rows(values_only=True)):
            if rnum > 3:
                break
            for cell in row:
                if cell and isinstance(cell, str) and "agency by job family" in cell.lower():
                    return ws
    return None

# ── Header parsing ────────────────────────────────────────────────────────────

# Cell values in the header area that are NOT job family names.
_NON_FAMILY = frozenset({"Agency", "Job Family", "Total", ".", ""})

_JF_FRAGMENTS = (
    "accounting", "administration", "business", "communications",
    "compliance", "data", "engineering", "health", "human resource",
    "ict", "information", "intelligence", "legal", "monitoring",
    "policy", "portfolio", "project", "science", "senior", "service",
    "trades", "no data", "zero data",
)


def find_jf_header_row(ws) -> tuple[int, dict[int, str]]:
    """Return (row_index_0based, {col_index: family_label}) for the JF header row."""
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 8:
            break
        candidates: dict[int, str] = {}
        for j, v in enumerate(row):
            if v is None:
                continue
            s = str(v).strip()
            if s and s not in _NON_FAMILY and j > 0:
                candidates[j] = s

        if len(candidates) < 5:
            continue

        combined = " ".join(candidates.values()).lower()
        hits = sum(1 for frag in _JF_FRAGMENTS if frag in combined)
        if hits >= 4:
            return i, candidates

    return -1, {}

# ── Value parsing ─────────────────────────────────────────────────────────────

def parse_headcount(v) -> int | None:
    """Return integer headcount, or None for cells shown as '.' in the source table.

    '.' most likely means zero employees or a not-applicable combination — stored as
    null rather than zero to preserve the source distinction.
    """
    if v is None:
        return None
    s = str(v).strip()
    if s in (".", "..", ""):
        return None
    try:
        return int(float(s.replace(",", "")))
    except ValueError:
        return None

# ── Sheet parsing ─────────────────────────────────────────────────────────────

def parse_sheet(ws, snapshot_date: str, regime: str) -> list[dict]:
    """Parse an agency × job family worksheet into long-format row dicts.

    Portfolio departments appear with no prefix; sub-agencies have a "- " prefix
    and are assigned parent_agency = the most recent portfolio department row.
    """
    header_row_idx, jf_cols = find_jf_header_row(ws)
    if not jf_cols:
        raise ValueError("Could not locate job family header row")

    records: list[dict] = []
    current_parent: str | None = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_row_idx:
            continue

        raw = row[0]
        if raw is None:
            continue
        agency_str = str(raw).strip()
        if not agency_str:
            continue

        lower = agency_str.lower()
        if lower == "total" or any(lower.startswith(p) for p in FOOTER_PREFIXES):
            continue

        is_sub = agency_str.startswith("- ")
        agency_clean = agency_str[2:].strip() if is_sub else agency_str

        if not is_sub:
            current_parent = agency_clean
            parent = None
        else:
            parent = current_parent

        for col_idx, family in jf_cols.items():
            cell_val = row[col_idx] if col_idx < len(row) else None
            family_norm = _FAMILY_NORMALISE.get(family, family)
            records.append({
                "snapshot_date":    snapshot_date,
                "taxonomy_regime":  regime,
                "agency":           agency_clean,
                "agency_canonical": _AGENCY_CANONICAL.get(agency_clean, agency_clean),
                "is_sub_agency":    is_sub,
                "parent_agency":    parent,
                "job_family":       family_norm,
                "job_family_key":   _JF_KEY.get(family_norm),
                "headcount":        parse_headcount(cell_val),
            })

    return records

# ── Main ──────────────────────────────────────────────────────────────────────

def main(target_date: str | None = None) -> None:
    files = sorted(XLSX_DIR.glob("apsed_*.xlsx"))
    all_records: list[dict] = []
    processed: list[str] = []

    for path in files:
        date = snapshot_date_from_path(path)
        if not date:
            continue
        if target_date and date != target_date:
            continue
        regime = REGIME_MAP.get(date)
        if not regime:
            print(f"  [warn] No regime mapping for {date}, skipping")
            continue

        print(f"  [{date}] Opening ...", end=" ", flush=True)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = find_agency_jf_sheet(wb)

        if ws is None:
            wb.close()
            print("no agency×JF table — skipped")
            continue

        try:
            records = parse_sheet(ws, date, regime)
        except Exception as exc:
            wb.close()
            print(f"ERROR: {exc}")
            traceback.print_exc()
            continue

        wb.close()

        n_agencies = len({r["agency"] for r in records if not r["is_sub_agency"]})
        n_families = len({r["job_family"] for r in records})
        print(f"{len(records):,} rows ({n_agencies} portfolio depts, {n_families} job families)")
        all_records.extend(records)
        processed.append(date)

    if not all_records:
        print("No rows parsed — nothing written.")
        return

    with open(OUTPUT, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLS)
        writer.writeheader()
        writer.writerows(all_records)

    parquet_path = OUTPUT.with_suffix(".parquet")
    arrays = []
    for col in OUTPUT_COLS:
        field_type = _SCHEMA.field(col).type
        values = [r[col] for r in all_records]
        if field_type == pa.date32():
            values = [datetime.date.fromisoformat(v) for v in values]
        arrays.append(pa.array(values, type=field_type))
    pq.write_table(pa.table(dict(zip(OUTPUT_COLS, arrays)), schema=_SCHEMA), parquet_path)

    print(f"\nWrote {len(all_records):,} rows → {OUTPUT}")
    print(f"                        → {parquet_path}")
    print(f"Dates: {processed}")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    main(target)
